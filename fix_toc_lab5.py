#!/usr/bin/env python3
"""Second pass for Lab5: its TOC lists both levels (lab titles + sections),
but the cached field result showed level-1 only. Rebuild the cached entries
with level-2 items included and indented; the field itself is unchanged."""
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, Cm

path = 'Lab5/Lab5_PlesuDinu_FAF-241.docx'
doc = Document(path)
paras = doc.paragraphs
ci = next(i for i, p in enumerate(paras) if p.text.strip() == 'Contents')

# current field paragraphs = those between Contents and the first Heading
field_pars = []
for p in paras[ci + 1:]:
    st = getattr(p.style, 'name', '') or ''
    if st.startswith('Heading'):
        break
    field_pars.append(p)

# heading inventory (bookmarks were added in the first pass)
heads = []
for p in paras[ci + 1:]:
    st = getattr(p.style, 'name', '') or ''
    if not st.startswith('Heading') or not p.text.strip():
        continue
    lvl = 2 if st.endswith('2') else 1
    bm = p._p.find(qn('w:bookmarkStart'))
    name = bm.get(qn('w:name')) if bm is not None else None
    heads.append((lvl, p.text.strip(), name))

def fld(par, kind, text=None):
    r = OxmlElement('w:r')
    if kind in ('begin', 'separate', 'end'):
        fc = OxmlElement('w:fldChar')
        fc.set(qn('w:fldCharType'), kind)
        r.append(fc)
    else:
        it = OxmlElement('w:instrText')
        it.set(qn('xml:space'), 'preserve')
        it.text = text
        r.append(it)
    par._p.append(r)

new_pars = []
for i, (lvl, text, name) in enumerate(heads):
    np = doc.add_paragraph()
    np.paragraph_format.space_after = Pt(4)
    if lvl == 2:
        np.paragraph_format.left_indent = Cm(0.75)
    if i == 0:
        fld(np, 'begin')
        fld(np, 'instr', ' TOC \\o "1-2" \\h \\z \\u ')
        fld(np, 'separate')
    np.add_run(text).font.size = Pt(12)
    if i == len(heads) - 1:
        fld(np, 'end')
    new_pars.append(np)

anchor = paras[ci]._p
for p in field_pars:
    p._p.getparent().remove(p._p)
for np in new_pars:
    anchor.addnext(np._p)
    anchor = np._p

doc.save(path)
print(f'{path}: cached TOC now lists {len(heads)} entries (both levels)')
