#!/usr/bin/env python3
"""Build Lab6 report docx from the UTM template using today's real xv6 captures."""
import os
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

TPL = '/home/user/Documents/GitHub/os-2026-PlesuDinu/resources/UTM_Lab_Report_Template.docx'
PNG = '/home/user/Documents/GitHub/os-2026-PlesuDinu/Lab6/captures/png'
OUT = '/home/user/Documents/GitHub/os-2026-PlesuDinu/Lab6/Lab6_PlesuDinu_FAF-241.docx'

doc = Document(TPL)

def set_text(p, text):
    for r in p.runs:
        r.text = ''
    if p.runs:
        p.runs[0].text = text
    else:
        p.add_run(text)

paras = doc.paragraphs
for p in paras:
    if '[Student Name]' in p.text:
        set_text(p, 'Pleșu Dinu FAF-241')
for p in paras:
    if '[N]' in p.text:
        set_text(p, p.text.replace('[N]', '6'))
    if '[Laboratory topic]' in p.text:
        set_text(p, p.text.replace('[Laboratory topic]', 'Inside the xv6 File System'))
    if '[Instructor, academic title]' in p.text:
        set_text(p, p.text.replace('[Instructor, academic title]', 'Reitman Patricia, university assistant'))

idx_contents = next(i for i, p in enumerate(paras) if p.text.strip() == 'Contents')
toc = ['1.  Purpose of the Work',
       '2.  Part 1 – Read what xv6 shows you',
       '3.  Part 2 – Write mystat',
       '4.  Part 3 – Measure the limits',
       '5.  Part 4 – Write find',
       '6.  Conclusion']
for p in paras[idx_contents + 1:]:
    p._p.getparent().remove(p._p)
for entry in toc:
    tp = doc.add_paragraph(entry)
    tp.paragraph_format.space_after = Pt(4)

def fresh_heading(text, level=1):
    h = doc.add_heading(text, level=level)
    for r in h.runs:
        r.font.color.rgb = RGBColor(0, 0, 0)
    return h

def body(text):
    p = doc.add_paragraph(text)
    p.paragraph_format.space_after = Pt(6)
    return p

def code(text):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Cm(1)
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run(text)
    r.font.name = 'Consolas'
    r.font.size = Pt(9.5)
    r.font.color.rgb = RGBColor(20, 20, 60)
    return p

def figure(png, caption):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(f'{PNG}/{png}.png', width=Cm(15.5))
    c = doc.add_paragraph()
    c.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = c.add_run(caption)
    r.italic = True
    r.font.size = Pt(10.5)

def obs(n, text):
    body(f'{n}. {text}')

doc.add_page_break()

# ==== 1. Purpose ====
fresh_heading('1.  Purpose of the Work')
body('The purpose of this laboratory is to read what the xv6 file system shows about its own '
     'objects, write my own stat program, measure the limits written in kernel/fs.h, and write '
     'find, which walks the directory tree recursively. Every session below was run on QEMU '
     'with the code in this folder (user/mystat.c, user/find.c and the Makefile with the '
     'UPROGS lines $U/_mystat and $U/_find).')
code('make clean ; make CPUS=1     # builds kernel + fs.img with mystat and find\n'
     'make qemu CPUS=1             # boot xv6; exit QEMU with Ctrl-A then X')
body('The build completed with exit status 0 on this machine, so both programs compile against '
     'the current xv6 tree.')

# ==== 2. Part 1 ====
fresh_heading('2.  Part 1 – Read what xv6 shows you')
body('Each ls line is: name, type, inode number, size in bytes. The type field is decoded '
     'from kernel/stat.h.')
figure('fig6-1-root-ls', 'Figure 1.1 – The root directory as xv6 shows it: name, type, inode, size')
code('mkdir a\nls a')
figure('fig6-2-empty-dir', 'Figure 1.2 – The empty directory a: two entries, size 32')

fresh_heading('2.1.  Observations', 2)
obs(1, 'In kernel/stat.h, type 1 is T_DIR (directory), type 2 is T_FILE (regular file) and '
       'type 3 is T_DEVICE (device file) — visible in Figure 1.1 as README = 2, console = 3.')
obs(2, 'The root directory has inode number 1 because kernel/fs.h defines ROOTINO as 1. '
       'In the root, . and .. are the same because the parent of the root is the root itself.')
obs(3, 'The empty directory a has size 32. One directory entry is 16 bytes: a 2-byte inode '
       'number plus a 14-byte name (DIRSIZ). The two entries are . (points to itself, inode 30) '
       'and .. (points to the parent, inode 1), as Figure 1.2 shows.')

# ==== 3. Part 2 ====
fresh_heading('3.  Part 2 – Write mystat')
body('user/mystat.c fills struct stat through the stat() library call and prints the name, '
     'type, inode, link count and size for every argument.')
code(open('/home/user/Documents/GitHub/os-2026-PlesuDinu/Lab6/mystat.c').read().rstrip())
figure('fig6-3-mystat', 'Figure 2.1 – mystat README . console on the fresh disk')
code('ln README r2\nmystat README r2')
figure('fig6-4-after-ln', 'Figure 2.2 – After ln: both names share inode 2, link count 2')
code('rm README\nmystat r2\ncat r2 | wc')
figure('fig6-5-after-rm', 'Figure 2.3 – After rm: the inode survives with link count 1; the data is unchanged')

fresh_heading('3.1.  Observations', 2)
obs(1, 'After ln, both README and r2 reported inode 2 with links 2: two names, one inode. '
       'After rm, r2 still reported inode 2 with size 2441 bytes, but links 1 — only the name '
       'was removed; what changed was the link count, what did not change were the inode number '
       'and the data (cat r2 | wc still printed the full content).')
obs(2, 'kernel/sysfile.c decrements the count inside sys_unlink with `ip->nlink--;` followed '
       'by `iupdate(ip);`. The data blocks are not freed at that moment: they are released only '
       'when the last link disappears and the inode is reclaimed (itrunc/bfree), so an open '
       'reader keeps working.')
obs(3, 'struct dinode in kernel/fs.h holds: short type, short major, short minor, short nlink, '
       'uint size and uint addrs[NDIRECT + 1] — type, device numbers, link count, length and '
       'block pointers. There is no symlink flag and no permission bits, which is why xv6 has '
       'neither symbolic links nor permissions.')

# ==== 4. Part 3 ====
fresh_heading('4.  Part 3 – Measure the limits')
code("grep -n \"define BSIZE\\|define NDIRECT\\|define NINDIRECT\\|define DIRSIZ\" kernel/fs.h\n"
     "make clean ; make fs.img   # mkfs prints the nmeta line")
figure('fig6-6-limits', 'Figure 3.1 – The constants in kernel/fs.h and the nmeta line of the build')
code('mkdir averyveryverylongname\nls')
figure('fig6-7-longname', 'Figure 3.2 – The long name is truncated to DIRSIZ = 14 characters')

fresh_heading('4.1.  Observations', 2)
obs(1, 'ls shows the directory as averyveryveryl: a directory entry keeps only DIRSIZ = 14 '
       'characters of the name, so the rest is cut off (Figure 3.2).')
obs(2, 'The largest file is 12 direct blocks plus one indirect block: '
       '12 × 1024 + (1024 / 4) × 1024 = 12 288 + 262 144 = 274 432 bytes, i.e. '
       '(NDIRECT + NINDIRECT) × BSIZE with NDIRECT = 12 and NINDIRECT = BSIZE / 4 = 256.')
obs(3, 'The nmeta line says: nmeta 47 (boot, super, log blocks 31, inode blocks 13, bitmap '
       'blocks 1) blocks 1953 total 2000. So the whole disk is 2000 blocks: 31 for the log, '
       '13 for inodes, 1 for the bitmap (plus boot and superblock make 47 metadata), and 1953 '
       'blocks are available for data.')

# ==== 5. Part 4 ====
fresh_heading('5.  Part 4 – Write find')
body('find <dir> <name> prints every path under <dir> whose last component is <name>. It opens '
     'a directory, reads struct dirent entries one by one, skips inum == 0, never recurses into '
     '. and .., compares names with strcmp, and closes every directory it opens (only 16 open '
     'files per process).')
code(open('/home/user/Documents/GitHub/os-2026-PlesuDinu/Lab6/find.c').read().rstrip())
body('Run on a fresh disk (make fs.img then make qemu), exactly as the assignment prescribes:')
figure('fig6-8-find', 'Figure 4.1 – find . r2 and find a b on a fresh disk')

fresh_heading('5.1.  Observations', 2)
obs(1, 'The walk found all three files named r2: ./a/b/r2, ./a/r2 and ./r2, in the order the '
       'directories are read. find a b printed a/b because only a/b matches the name b.')
body('The recursion is the same idea behind real filesystem tools: open, read entries, build a '
     'path, stat it, recurse, close — respecting the per-process open-file limit.')

# ==== 6. Conclusion ====
fresh_heading('6.  Conclusion')
body('The xv6 file system is small but instructive: directories are just files containing 16-byte '
     'dirent records, names are capped at DIRSIZ = 14, and inodes carry the metadata that manages '
     'size, hard links and block addresses (12 direct + 256 indirect = 274 432 bytes maximum). '
     'A hard link shares one inode; data blocks are reclaimed only when the last link disappears, '
     'which sys_unlink implements with a single nlink decrement. Writing mystat and find showed '
     'how the library calls map onto the kernel syscalls, and running both inside QEMU confirmed '
     'they compile and behave on the current xv6 tree.')

os.makedirs(os.path.dirname(OUT), exist_ok=True)
doc.save(OUT)
print('saved', OUT)
